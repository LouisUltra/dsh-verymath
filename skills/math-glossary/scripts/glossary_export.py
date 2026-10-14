"""术语库 Excel 导入导出工具。"""

from __future__ import annotations

import argparse
import json
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any


REQUIRED_FIELDS = ("term_id", "en", "cn")
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _load_terms(glossary_path: str | os.PathLike[str]) -> list[dict[str, Any]]:
    path = Path(glossary_path)
    with path.open("r", encoding="utf-8") as handle:
        terms = json.load(handle)
    if not isinstance(terms, list):
        raise ValueError("glossary.json 顶层必须是术语对象数组")
    for index, term in enumerate(terms, 1):
        if not isinstance(term, dict):
            raise ValueError(f"第 {index} 条术语必须是 JSON 对象")
        missing = [field for field in REQUIRED_FIELDS if not term.get(field)]
        if missing:
            raise ValueError(
                f"第 {index} 条术语缺少必填字段: {', '.join(missing)}"
            )
    return terms


def _atomic_json_dump(
    terms: list[dict[str, Any]],
    output_path: str | os.PathLike[str],
) -> None:
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.name}.",
        suffix=".tmp",
        dir=destination.parent,
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            json.dump(terms, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary_name, destination)
    except Exception:
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError(
            "缺少 openpyxl。请先审阅依赖，再安装 requirements.txt。"
        ) from exc
    return openpyxl


def _excel_safe(value: Any) -> Any:
    """Prevent user-controlled text from becoming an Excel formula."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if (
        len(text) > 1
        and text.startswith("'")
        and text[1:].startswith(FORMULA_PREFIXES)
    ):
        return text[1:]
    return text


def export_to_excel(
    glossary_path: str | os.PathLike[str],
    output_dir: str | os.PathLike[str],
) -> str:
    """将 glossary.json 原子导出为 Excel 文件，供人工审阅。"""
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill

    terms = _load_terms(glossary_path)
    destination_dir = Path(output_dir)
    destination_dir.mkdir(parents=True, exist_ok=True)
    output_path = destination_dir / (
        f"glossary_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    )

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "术语库"

    headers = [
        "序号",
        "英文术语",
        "中文译名",
        "词性",
        "子领域",
        "来源",
        "例句(英文)",
        "例句(中文)",
        "创建时间",
        "更新时间",
    ]
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid",
    )

    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(row=1, column=column, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_index, term in enumerate(terms, 2):
        examples = term.get("examples", [])
        if not isinstance(examples, list):
            raise ValueError(f"{term['term_id']} 的 examples 必须是数组")
        examples_en = "; ".join(
            example.get("en", "")
            for example in examples
            if isinstance(example, dict)
        )
        examples_cn = "; ".join(
            example.get("cn", "")
            for example in examples
            if isinstance(example, dict)
        )
        field_value = term.get("field", "")
        if isinstance(field_value, list):
            field_value = " / ".join(str(item) for item in field_value)

        values = [
            term.get("idx", row_index - 1),
            term.get("en", ""),
            term.get("cn", ""),
            term.get("pos", ""),
            field_value,
            term.get("source", ""),
            examples_en,
            examples_cn,
            term.get("created_at", ""),
            term.get("updated_at", ""),
        ]
        for column, value in enumerate(values, 1):
            worksheet.cell(
                row=row_index,
                column=column,
                value=_excel_safe(value),
            )

    for column, width in enumerate(
        [6, 22, 18, 10, 12, 20, 30, 30, 12, 12],
        1,
    ):
        worksheet.column_dimensions[
            openpyxl.utils.get_column_letter(column)
        ].width = width

    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.stem}.",
        suffix=".xlsx",
        dir=destination_dir,
    )
    os.close(descriptor)
    try:
        workbook.save(temporary_name)
        workbook.close()
        os.replace(temporary_name, output_path)
    except Exception:
        workbook.close()
        try:
            os.unlink(temporary_name)
        except FileNotFoundError:
            pass
        raise
    return str(output_path)


def _parse_field(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [
        part.strip()
        for part in _cell_text(value).split("/")
        if part.strip()
    ]


def import_from_excel(
    excel_path: str | os.PathLike[str],
    output_json_path: str | os.PathLike[str],
) -> list[dict[str, Any]]:
    """从 Excel 导入术语并原子写入 JSON。"""
    openpyxl = _openpyxl()
    workbook = openpyxl.load_workbook(excel_path, read_only=True)
    worksheet = workbook.active
    terms: list[dict[str, Any]] = []
    today = datetime.now().strftime("%Y-%m-%d")

    try:
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 3 or not row[1]:
                continue

            en_value = _cell_text(row[1])
            cn_value = _cell_text(row[2])
            if not en_value or not cn_value:
                continue

            examples_en_raw = (
                _cell_text(row[6]) if len(row) > 6 else ""
            )
            examples_cn_raw = (
                _cell_text(row[7]) if len(row) > 7 else ""
            )
            english_examples = [
                item.strip()
                for item in examples_en_raw.split(";")
                if item.strip()
            ]
            chinese_examples = [
                item.strip()
                for item in examples_cn_raw.split(";")
                if item.strip()
            ]
            examples = [
                {
                    "en": example,
                    "cn": (
                        chinese_examples[index]
                        if index < len(chinese_examples)
                        else ""
                    ),
                }
                for index, example in enumerate(english_examples)
            ]

            terms.append(
                {
                    "term_id": "",
                    "en": en_value,
                    "cn": cn_value,
                    "pos": (
                        _cell_text(row[3]) if len(row) > 3 else ""
                    ),
                    "field": _parse_field(
                        row[4] if len(row) > 4 else None
                    ),
                    "source": (
                        _cell_text(row[5]) if len(row) > 5 else ""
                    ),
                    "examples": examples,
                    "created_at": (
                        _cell_text(row[8])
                        if len(row) > 8 and row[8]
                        else today
                    ),
                    "updated_at": today,
                }
            )
    finally:
        workbook.close()

    terms.sort(
        key=lambda term: (
            term.get("en", "").lower(),
            term.get("cn", "").lower(),
        )
    )
    occurrences: dict[str, int] = {}
    for index, term in enumerate(terms, 1):
        base_id = term["en"].lower()
        occurrence = occurrences.get(base_id, 0)
        term["term_id"] = base_id if occurrence == 0 else f"{base_id}__{occurrence}"
        occurrences[base_id] = occurrence + 1
        term["idx"] = index

    _atomic_json_dump(terms, output_json_path)
    return terms


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="在 glossary.json 与审阅用 Excel 之间转换。"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    export_parser = subparsers.add_parser("export", help="JSON 导出为 Excel")
    export_parser.add_argument("glossary", help="glossary.json 路径")
    export_parser.add_argument("output_dir", help="Excel 输出目录")

    import_parser = subparsers.add_parser("import", help="Excel 导入为 JSON")
    import_parser.add_argument("excel", help="输入 .xlsx 路径")
    import_parser.add_argument("output_json", help="输出 glossary.json 路径")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.command == "export":
        result = export_to_excel(args.glossary, args.output_dir)
        print(f"Exported to: {result}")
    else:
        terms = import_from_excel(args.excel, args.output_json)
        print(f"Imported {len(terms)} terms to: {args.output_json}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
